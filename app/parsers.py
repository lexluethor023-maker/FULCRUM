"""Read-only extraction. Originals remain in the content-addressed evidence store."""
import csv
import io
import json
import re
import zipfile
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
from xml.etree import ElementTree

PARSER_VERSION = 'corpus-1'
SUPPORTED = {'.txt', '.md', '.csv', '.tsv', '.json', '.jsonl', '.html', '.htm', '.pdf', '.docx', '.xlsx'}
MAX_UNITS = 250000
MAX_TEXT = 200_000_000


@dataclass
class Unit:
    locator: str
    text: str
    fields: dict = field(default_factory=dict)


def value_text(value):
    if value is None:
        return ''
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, default=str)
    return str(value)


def record_unit(locator, record):
    fields = {str(k): value_text(v) for k, v in record.items()}
    return Unit(locator, '\n'.join(f'{k}: {v}' for k, v in fields.items() if v.strip()), fields)


def headers_unique(values):
    seen = set()
    names = []
    for i, value in enumerate(values, 1):
        base = value_text(value).strip() or f'column_{i}'
        name = base
        suffix = 2
        while name in seen:
            name = f'{base}_{suffix}'
            suffix += 1
        seen.add(name)
        names.append(name)
    return names


def decode(raw):
    if raw.startswith((b'\xff\xfe', b'\xfe\xff')):
        return raw.decode('utf-16')
    return raw.decode('utf-8-sig')


def text_units(text, prefix='text'):
    # Character offsets refer to this extracted text, not byte offsets in the original.
    for offset in range(0, len(text), 6000):
        chunk = text[offset:offset + 6000]
        if chunk.strip():
            yield Unit(f'{prefix}, characters {offset + 1}-{offset + len(chunk)}', chunk)


class HTMLText(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.hidden = 0
        self.parts = []
        self.links = []

    def handle_starttag(self, tag, attrs):
        if tag in ('script', 'style', 'noscript', 'template'):
            self.hidden += 1
        if not self.hidden:
            if tag in ('p', 'div', 'br', 'li', 'h1', 'h2', 'h3', 'tr'):
                self.parts.append('\n')
            if tag == 'a':
                href = dict(attrs).get('href', '')
                if href.startswith(('https://', 'http://')):
                    self.links.append(href)

    def handle_endtag(self, tag):
        if tag in ('script', 'style', 'noscript', 'template') and self.hidden:
            self.hidden -= 1

    def handle_data(self, data):
        if not self.hidden:
            self.parts.append(data)


def check_zip(path):
    with zipfile.ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 50000 or sum(x.file_size for x in members) > 512 * 1024 * 1024:
            raise ValueError('Expanded Office document exceeds parser limits.')


def extract(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix not in SUPPORTED:
        raise ValueError(f'Unsupported format: {suffix}')
    count = size = 0
    for unit in _extract(path, suffix):
        if not unit.text.strip():
            continue
        count += 1
        size += len(unit.text)
        if count > MAX_UNITS or size > MAX_TEXT:
            raise ValueError('Document exceeds extraction limits; split into smaller files.')
        yield unit
    if not count:
        raise ValueError('No readable text or records; scanned documents require OCR.')


def _extract(path, suffix):
    if suffix in ('.txt', '.md'):
        yield from text_units(decode(path.read_bytes()))
    elif suffix in ('.html', '.htm'):
        parser = HTMLText()
        parser.feed(decode(path.read_bytes()))
        yield from text_units(''.join(parser.parts), 'HTML text')
        if parser.links:
            yield record_unit('HTML outbound links', {'source_links': '\n'.join(dict.fromkeys(parser.links))})
    elif suffix in ('.csv', '.tsv'):
        csv.field_size_limit(4 * 1024 * 1024)
        text = decode(path.read_bytes())
        delimiter = '\t' if suffix == '.tsv' else ','
        reader = csv.reader(io.StringIO(text), delimiter=delimiter)
        headers = None
        for row_number, row in enumerate(reader, 1):
            if not any(x.strip() for x in row):
                continue
            if headers is None:
                headers = headers_unique(row)
            else:
                expanded = headers + [f'extra_column_{i}' for i in range(len(headers) + 1, len(row) + 1)]
                yield record_unit(f'row {row_number}', dict(zip(expanded, row)))
    elif suffix in ('.json', '.jsonl'):
        text = decode(path.read_bytes())
        records = ((f'line {i}', json.loads(line)) for i, line in enumerate(text.splitlines(), 1) if line.strip()) if suffix == '.jsonl' else [("$", json.loads(text))]
        for locator, value in records:
            rows = enumerate(value) if isinstance(value, list) else [(None, value)]
            for index, row in rows:
                loc = f'{locator}[{index}]' if index is not None else locator
                yield record_unit(loc, row if isinstance(row, dict) else {'value': row})
    elif suffix == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('Encrypted PDF requires an accessible source copy.')
        if len(reader.pages) > 10000:
            raise ValueError('PDF exceeds 10,000-page limit.')
        for i, page in enumerate(reader.pages, 1):
            yield from text_units(page.extract_text() or '', f'PDF page {i}')
    elif suffix == '.docx':
        check_zip(path)
        with zipfile.ZipFile(path) as archive:
            xml = ElementTree.fromstring(archive.read('word/document.xml'))
            namespace = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
            for i, paragraph in enumerate(xml.iter(namespace + 'p'), 1):
                text = ''.join(node.text or '' for node in paragraph.iter(namespace + 't'))
                if text.strip():
                    yield from text_units(text, f'DOCX paragraph {i}')
    elif suffix == '.xlsx':
        from openpyxl import load_workbook
        check_zip(path)
        workbook = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        cached = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        try:
            for sheet in workbook:
                if sheet.max_column and sheet.max_column > 512:
                    raise ValueError(f'Sheet {sheet.title} exceeds 512-column limit.')
                if sheet.max_row and sheet.max_row > 1000000:
                    raise ValueError(f'Sheet {sheet.title} exceeds 1,000,000-row limit.')
                headers = None
                for row_number, (cells, cached_cells) in enumerate(zip(sheet.iter_rows(), cached[sheet.title].iter_rows()), 1):
                    row = [cell.value for cell in cells]
                    if not any(value_text(x).strip() for x in row):
                        continue
                    if headers is None:
                        headers = headers_unique(row)
                        continue
                    fields = dict(zip(headers, row))
                    for name, cell, saved in zip(headers, cells, cached_cells):
                        if getattr(cell, 'data_type', None) == 'f':
                            fields[name + '__formula'] = cell.value
                            fields[name] = saved.value if saved.value is not None else cell.value
                    yield record_unit(f'{sheet.title}!row {row_number}', fields)
        finally:
            workbook.close()
            cached.close()
